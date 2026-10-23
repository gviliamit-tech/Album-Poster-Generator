import argparse
import json
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path
import requests

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt
from sklearn.cluster import KMeans

# ----------------------------------------------------------------------------
# CONFIG 
# ----------------------------------------------------------------------------
FONT_MAIN_DEFAULT = "Bahnschrift SemiBold Condensed"
FONT_HEBREW_DEFAULT = "Arial"  # פונט ברירת מחדל התומך בעברית
FONT_LIGHT = "Bahnschrift Light"
BOLD = False
COND_EM = 0.50

PAGE_W, PAGE_H = 8.2677, 11.6929   # A4 in inches
MARGIN = 0.80
COVER_Y = 0.35
RULE_Y = 10.62
LIST_BOTTOM = 10.40
GRAY_TEXT = "808080"

# בדיקה האם יש טקסט בעברית
def contains_hebrew(text):
    return bool(re.search(r'[\u0590-\u05FF]', str(text)))

# ============================== COLOR HELPERS ===============================
def hex_to_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

def rgb_to_hex(rgb):
    return "".join(f"{int(round(max(0, min(255, c)))):02X}" for c in rgb)

def rgb_to_lab(rgb):
    rgb = np.where(rgb > 0.04045, ((rgb + 0.055) / 1.055) ** 2.4, rgb / 12.92)
    m = np.array([[0.4124564, 0.3575761, 0.1804375],
                  [0.2126729, 0.7151522, 0.0721750],
                  [0.0193339, 0.1191920, 0.9503041]])
    xyz = (rgb @ m.T) / np.array([0.95047, 1.0, 1.08883])
    f = np.where(xyz > 0.008856, np.cbrt(xyz), 7.787 * xyz + 16 / 116)
    return np.stack([116 * f[:, 1] - 16, 500 * (f[:, 0] - f[:, 1]), 200 * (f[:, 1] - f[:, 2])], axis=1)

def luminance(hexcolor):
    def ch(c):
        c /= 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(c) for c in hex_to_rgb(hexcolor))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b

def contrast_on_white(hexcolor):
    return 1.05 / (luminance(hexcolor) + 0.05)

def saturation(hexcolor):
    r, g, b = (c / 255 for c in hex_to_rgb(hexcolor))
    mx, mn = max(r, g, b), min(r, g, b)
    return 0 if mx == 0 else (mx - mn) / mx

def extract_palette(img, n=5, k=16):
    im = img.convert("RGB")
    im.thumbnail((140, 140))
    arr = np.asarray(im).reshape(-1, 3).astype(float) / 255
    lab = rgb_to_lab(arr)
    k = min(k, len(np.unique(arr, axis=0)))
    km = KMeans(n_clusters=k, n_init=4, random_state=0).fit(lab)

    cands = []
    for i in range(k):
        mask = km.labels_ == i
        share = mask.mean()
        if share < 0.004:
            continue
        cands.append({"share": share, "lab": km.cluster_centers_[i],
                      "rgb": arr[mask].mean(axis=0) * 255})

    def weight(c):
        chroma = float(np.hypot(c["lab"][1], c["lab"][2]))
        return (c["share"] ** 0.4) * (1 + 1.5 * min(chroma, 60) / 60)

    chosen = [max(range(len(cands)), key=lambda i: cands[i]["share"])]
    for min_de in (22, 14, 8, 0):
        while len(chosen) < n:
            best, best_score = None, 0
            for i, c in enumerate(cands):
                if i in chosen:
                    continue
                d = min(np.linalg.norm(c["lab"] - cands[j]["lab"]) for j in chosen)
                if d < min_de:
                    continue
                if d * weight(c) > best_score:
                    best, best_score = i, d * weight(c)
            if best is None:
                break
            chosen.append(best)
        if len(chosen) >= n:
            break
    while len(chosen) < n:
        chosen.append(chosen[len(chosen) % len(chosen)])

    colors = [rgb_to_hex(cands[i]["rgb"]) for i in chosen]
    return sorted(colors, key=luminance, reverse=True)

def pick_text_colors(palette):
    readable = sorted([c for c in palette if contrast_on_white(c) >= 3.0], key=luminance)
    if not readable:
        return "222222", "555555"
    title = readable[1] if len(readable) > 2 and contrast_on_white(readable[1]) >= 7 else readable[0]
    rest = [c for c in readable if c != title and contrast_on_white(c) <= 9] or readable[-1:]
    artist = max(rest, key=saturation)
    return title, artist

# ============================== TEXT HELPERS ================================
def est_width(text, pt, em=None):
    return len(text) * pt / 72 * (em or COND_EM)

def split_balanced(text):
    words = text.split()
    if len(words) < 2:
        return [text]
    best = min(range(1, len(words)),
               key=lambda i: abs(len(" ".join(words[:i])) - len(" ".join(words[i:]))))
    return [" ".join(words[:best]), " ".join(words[best:])]

def add_text(slide, name, x, y, w, h, lines, font, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_pt=None, spacing_pt=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_pt:
            p.line_spacing = Pt(line_pt)
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        if spacing_pt:
            r._r.get_or_add_rPr().set("spc", str(int(spacing_pt * 100)))
    return tb

def add_rect(slide, name, x, y, w, h, hexcolor):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(hexcolor)
    s.line.fill.background()
    s.shadow.inherit = False
    style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        s._element.remove(style)
    return s

def fmt_date(s):
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m", "%Y"):
        try:
            d = datetime.strptime(s, fmt)
        except ValueError:
            continue
        if fmt == "%Y-%m-%d":
            return d.strftime("%B %d, %Y").upper().replace(" 0", " ")
        if fmt == "%Y-%m":
            return d.strftime("%B %Y").upper()
        return s
    return s.upper()

# ============================== API AUTOMATION ==============================
def fetch_album_data(search_query):
    print(f"Searching internet for: '{search_query}'...")
    url = "https://itunes.apple.com/search"
    response = requests.get(url, params={"term": search_query, "entity": "album", "limit": 10}).json()
    
    if not response['results']:
        sys.exit("Album not found! Try checking your spelling.")
    
    best_match = response['results'][0]
    
    if " - " in search_query:
        target_artist = search_query.split(" - ")[0].strip().lower()
        for res in response['results']:
            if target_artist in res.get('artistName', '').lower():
                best_match = res
                break
                
    album_id = best_match['collectionId']
    cover_url = best_match['artworkUrl100'].replace("100x100bb", "1000x1000bb")
    title = best_match['collectionName']
    artist = best_match['artistName']
    date = best_match.get('releaseDate', '')[:10]
    genre = best_match.get('primaryGenreName', '').replace('/', ' | ')

    print(f"Found Match: {artist} - {title}")

    track_url = "https://itunes.apple.com/lookup"
    track_response = requests.get(track_url, params={"id": album_id, "entity": "song"}).json()
    
    tracks = []
    for item in track_response['results']:
        if item['wrapperType'] == 'track':
            tracks.append(item['trackName'])
            
    print("Downloading high-res cover art...")
    img_data = requests.get(cover_url).content
    temp_cover_path = Path(tempfile.gettempdir()) / "temp_cover.jpg"
    with open(temp_cover_path, "wb") as f:
        f.write(img_data)
        
    return {
        "artist": artist,
        "title": title,
        "date": date,
        "genres": genre,
        "tracks": tracks,
        "cover_path": str(temp_cover_path)
    }

# ============================== BUILD THE SLIDE =============================
# ============================== BUILD THE SLIDE =============================
def build_poster(cover_path, artist, title, date, genres, tracks, out_path):
    all_text_check = f"{artist} {title} {' '.join(tracks)}"
    is_rtl = contains_hebrew(all_text_check)
    
    font_main = FONT_HEBREW_DEFAULT if is_rtl else FONT_MAIN_DEFAULT

    artist_u = artist.upper() if not contains_hebrew(artist) else artist
    title_u = title.upper() if not contains_hebrew(title) else title
    
    tracks_u = []
    for i, t in enumerate(tracks):
        clean_t = re.sub(r"^\s*\d+[.)]\s*", "", t).strip()
        clean_t = clean_t.upper() if not contains_hebrew(clean_t) else clean_t
        # Format track number to the right of text if RTL
        tracks_u.append(f"{i + 1}. {clean_t}" if not is_rtl else f"{clean_t} .{i + 1}")

    img = Image.open(cover_path)
    palette = extract_palette(img)
    title_color, artist_color = pick_text_colors(palette)

    pic_path = cover_path
    if img.format not in ("JPEG", "PNG"):
        pic_path = Path(tempfile.mkdtemp()) / "cover.png"
        img.convert("RGB").save(pic_path)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) 

    cw = PAGE_W - 2 * MARGIN
    right = MARGIN + cw
    s = cw / 1349
    font_kw = dict(font=font_main, bold=BOLD)

    # --- Cover ---
    pic = slide.shapes.add_picture(str(pic_path), Inches(MARGIN), Inches(COVER_Y), Inches(cw), Inches(cw))
    pic.name = "Cover"
    w, h = img.size
    if w > h:
        pic.crop_left = pic.crop_right = (w - h) / 2 / w
    elif h > w:
        pic.crop_top = pic.crop_bottom = (h - w) / 2 / h
    cover_bottom = COVER_Y + cw

    # --- Swatches (RTL pushes them to the LEFT) ---
    sw_y = cover_bottom + 0.21
    sw_total, gap, sw_h = 585 * s, 25 * s, 91 * s
    sw_w = (sw_total - 4 * gap) / 5
    sw_start_x = MARGIN if is_rtl else right - sw_total
    
    for i, col in enumerate(palette):
        add_rect(slide, f"Swatch {i + 1}", sw_start_x + i * (sw_w + gap), sw_y, sw_w, sw_h, col)

    # --- Genres (RTL pushes them to the LEFT) ---
    if genres:
        g_x = MARGIN if is_rtl else right - 3.2
        g_align = PP_ALIGN.LEFT if is_rtl else PP_ALIGN.RIGHT
        add_text(slide, "Genres", g_x, sw_y + sw_h + 0.10, 3.2, 0.3, [genres],
                 FONT_LIGHT if not is_rtl else FONT_HEBREW_DEFAULT, 11.5, GRAY_TEXT, align=g_align)

    # --- Tracklist (RTL pushes it to the RIGHT) ---
    list_top = cover_bottom + 0.25
    list_w = 0.44 * cw
    list_align = PP_ALIGN.RIGHT if is_rtl else PP_ALIGN.LEFT
    list_x = right - list_w if is_rtl else MARGIN
    
    FONT_SIZE_PT = 14
    LINE_SPACING_PT = 16.8
    
    if len(tracks_u) > 13:
        col1 = tracks_u[:13]
        col2 = tracks_u[13:]
        col_w = list_w * 0.60
        
        # In RTL, Column 1 is on the far right, Column 2 is to its left
        x1 = right - col_w if is_rtl else MARGIN
        x2 = right - col_w * 2 if is_rtl else MARGIN + col_w
        
        add_text(slide, "Tracklist 1", x1, list_top, col_w, LIST_BOTTOM - list_top,
                 col1, color="111111", size=FONT_SIZE_PT, line_pt=LINE_SPACING_PT, align=list_align, **font_kw)
        add_text(slide, "Tracklist 2", x2, list_top, col_w, LIST_BOTTOM - list_top,
                 col2, color="111111", size=FONT_SIZE_PT, line_pt=LINE_SPACING_PT, align=list_align, **font_kw)
    else:
        add_text(slide, "Tracklist", list_x, list_top, list_w, LIST_BOTTOM - list_top,
                 tracks_u, color="111111", size=FONT_SIZE_PT, line_pt=LINE_SPACING_PT, align=list_align, **font_kw)

    add_rect(slide, "Rule", MARGIN - 0.12, RULE_Y, cw + 0.24, 0.05, "000000")
    
    # --- Date (RTL pushes it to the RIGHT) ---
    date_align = PP_ALIGN.RIGHT if is_rtl else PP_ALIGN.LEFT
    date_x = right - 4 if is_rtl else MARGIN
    add_text(slide, "Date", date_x, RULE_Y + 0.12, 4, 0.3, [fmt_date(date)], color="111111", size=14, align=date_align, **font_kw)

    # --- Artist + Title Block (RTL pushes it to the LEFT) ---
    max_block_w = 0.54 * cw
    TITLE_EM = COND_EM * 0.96
    t_align = PP_ALIGN.LEFT if is_rtl else PP_ALIGN.RIGHT
    
    t_lines = [title_u]
    pt = min(54, max_block_w * 72 / (len(title_u) * TITLE_EM))
    
    if pt < 28 and len(title_u.split()) > 1:
        t_lines = split_balanced(title_u)
        pt = min(46, max_block_w * 72 / (max(len(l) for l in t_lines) * TITLE_EM))
        
    pt = max(pt * 0.95, 14)
    t_h = len(t_lines) * pt * 1.0 / 72
    descent = lambda pt: 0.20 * pt / 72
    t_bottom = RULE_Y - 0.06 + descent(pt)
    
    t_x = MARGIN if is_rtl else right + 0.10 - max_block_w
    
    # Title
    add_text(slide, "Album title", t_x, t_bottom - t_h - 0.1, max_block_w, t_h + 0.1,
             t_lines, color=title_color, size=pt, align=t_align, anchor=MSO_ANCHOR.BOTTOM,
             spacing_pt=1, line_pt=pt, **font_kw)
             
    # Artist
    a_pt = min(22, max_block_w * 72 / (len(artist_u) * COND_EM))
    add_text(slide, "Artist", t_x, t_bottom - t_h - 0.1 - a_pt * 1.2 / 72, max_block_w,
             a_pt * 1.2 / 72, [artist_u], color=artist_color, size=a_pt, align=t_align,
             anchor=MSO_ANCHOR.BOTTOM, **font_kw)

    prs.save(out_path)
    return palette

# ============================== INPUT HANDLING ==============================
def main():
    ap = argparse.ArgumentParser(description="Make an editable A4 album poster (.pptx)")
    ap.add_argument("--search", help="Search for an album online")
    a, _ = ap.parse_known_args()

    if not a.search:
        sys.exit("You must provide an album to search! Example: python make_poster.py --search \"טונה - טונהפארק\"")

    meta = fetch_album_data(a.search)
    
    clean_filename = re.sub(r'[\\/:*?"<>|]', "", f"{meta['artist']} - {meta['title']}") + ".pptx"
    out_dir = Path("posters")
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / clean_filename

    print("Generating poster layout...")
    palette = build_poster(meta["cover_path"], meta["artist"], meta["title"], meta.get("date", ""), meta.get("genres", ""), meta.get("tracks", []), out)
    print(f"Success! Saved poster to: {out}")
    print("Color Palette used:", " ".join("#" + c for c in palette))

if __name__ == "__main__":
    main()